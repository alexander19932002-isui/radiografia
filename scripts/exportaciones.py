
"""Exportaciones del panel. Los filtros recibidos definen el informe."""
from io import BytesIO
from datetime import datetime
from html import escape
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import reportlab
from pathlib import Path
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table as PDFTable, TableStyle, PageBreak
from scripts.almacenamiento import comparar

INDICADORES = {
    'colocaciones_mmclp':'Colocaciones (MM CLP)',
    'morosidad_90_pct':'Morosidad 90+ (%)',
    'provisiones_clasificacion_pct':'Provisiones clasificación (%)',
    'eficiencia_gastos_ingresos_pct':'Gastos / ingresos (%)',
    'roe_pct':'ROE anualizado (%)',
    'roa_pct':'ROA anualizado (%)',
    'utilidad_acumulada_mmclp':'Utilidad acumulada (MM CLP)',
    'participacion_colocaciones_pct':'Participación colocaciones (%)',
}
FUENTE = 'https://www.cmfchile.cl/portal/estadisticas/626/w4-propertyvalue-28910.html'
NOTAS = [
    'ROE y ROA oficiales: después de impuestos y anualizados, sobre patrimonio y activos al cierre respectivamente. La utilidad/pérdida consolidada es acumulada desde enero; no es el flujo de un mes ni el resultado atribuible exclusivamente a propietarios.',
    'Participación en colocaciones = colocaciones totales del banco / colocaciones totales oficiales del sistema x 100, con igual fecha y alcance consolidado. El filtro de bancos no cambia el denominador. Para el sistema corresponde 100%.',
    'Fuente: reportes mensuales públicos de la CMF. Alcance consolidado. Los montos están en millones de pesos chilenos (MM CLP).',
    'Morosidad: créditos con 90 días o más de atraso / colocaciones totales. Provisiones: índice por clasificación sobre colocaciones a costo amortizado; no es un monto ni incluye todas las variantes adicionales.',
    'Eficiencia: gastos / ingresos operacionales acumulados al cierre. Comparar dos cierres no equivale a medir el gasto aislado de cada mes.',
    'Las diferencias de tasas se expresan en puntos porcentuales. Las variaciones de montos son nominales, sin ajuste por inflación. No se calculan variaciones con base cero ni variaciones porcentuales de utilidades con base negativa. Sí se muestran las diferencias monetarias.',
    'El sistema bancario se presenta aparte y no se suma a los bancos. No se promedian tasas para reconstruir el indicador oficial del sistema.',
    'Banco Bice absorbió Banco Security el 1 de noviembre de 2025, según las notas de los reportes. Considerar el cambio al comparar períodos anteriores.',
]

def numero(valor, dec=2):
    return 'No disponible' if pd.isna(valor) else f'{valor:,.{dec}f}'.translate(str.maketrans({',':'.', '.':','}))

def excel(tabla, mes, referencia, indicador, bancos):
    """Genera un XLSX con números tipados y filtros, sin ejecutar texto como fórmulas."""
    wb = Workbook()
    wb.remove(wb.active)
    def hoja(nombre, df):
        ws = wb.create_sheet(nombre)
        ws.append(list(df.columns))
        for fila in df.itertuples(index=False, name=None):
            ws.append([None if pd.isna(v) else v for v in fila])
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='17324D')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[1].height=42
        ws.freeze_panes='B2'
        for columna in ws.columns:
            titulo=columna[0].value
            ws.column_dimensions[columna[0].column_letter].width = 43 if titulo in ('banco','archivo_fuente','nota') else 24
            for cell in columna[1:]:
                if isinstance(cell.value,str):
                    cell.data_type='s'
                elif isinstance(cell.value,(int,float)):
                    cell.number_format='#,##0.00'
        if len(df):
            t=Table(displayName='Tabla'+str(len(wb.worksheets)),ref=ws.dimensions)
            t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
            ws.add_table(t)
        return ws
    seleccion=tabla[(tabla.banco.isin(bancos)) & (tabla.tipo_entidad=='banco')]
    hoja('Bancos del mes',seleccion[seleccion.fecha==mes])
    if referencia:
        hoja('Cambios',comparar(tabla,mes,referencia,indicador,bancos))
    hoja('Historial seleccionado',seleccion.sort_values(['fecha','banco']))
    hoja('Sistema oficial',tabla[tabla.tipo_entidad=='sistema'].sort_values('fecha'))
    notas=pd.DataFrame({'nota':[
        f'Mes analizado: {mes}. Referencia: {referencia or "sin referencia"}. Indicador de cambios: {indicador}.',
        'Los campos con sufijo _pct están en escala 0-100: 2.33 significa 2.33%. No aplicar directamente el formato porcentual de Excel.',
        'En Cambios, actual y referencia usan la unidad del indicador. cambio usa MM CLP para colocaciones y puntos porcentuales para tasas. variacion_pct es una tasa nominal en escala 0-100.',
        'El historial contiene todos los meses guardados de los bancos seleccionados; el sistema oficial contiene todas las fechas guardadas.',
        'Se conservan las fechas de cierre como claves ISO AAAA-MM-DD para facilitar uniones con SQL y Power BI.',
        *NOTAS,FUENTE]})
    ws=hoja('Definiciones y fuentes',notas)
    ws.column_dimensions['A'].width=120
    for row in ws.iter_rows(min_row=2):
        row[0].alignment=Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[row[0].row].height=48
    buf=BytesIO(); wb.save(buf); return buf.getvalue()

def pdf(tabla, mes, referencia, indicador, bancos):
    fuentes=Path(reportlab.__file__).parent/'fonts'
    if 'Informe' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('Informe',str(fuentes/'Vera.ttf')))
        pdfmetrics.registerFont(TTFont('Informe-Bold',str(fuentes/'VeraBd.ttf')))
        pdfmetrics.registerFontFamily('Informe',normal='Informe',bold='Informe-Bold',italic='Informe',boldItalic='Informe-Bold')
    buf=BytesIO()
    doc=SimpleDocTemplate(buf,pagesize=landscape(A4),rightMargin=30,leftMargin=30,topMargin=30,bottomMargin=35,title='Radiografía del sistema bancario chileno',author='Radiografía bancaria')
    estilos=getSampleStyleSheet()
    for estilo in estilos.byName.values():
        estilo.fontName='Informe-Bold' if estilo.name.startswith('Heading') or estilo.name=='Title' else 'Informe'
    estilos.add(ParagraphStyle(name='CeldaBanco',fontName='Informe',fontSize=8,leading=10))
    contenido=[Paragraph('Radiografía del sistema bancario chileno',estilos['Title']),Paragraph(f'Cierre: {mes} | Referencia: {referencia or "sin referencia"}',estilos['Normal']),Spacer(1,12)]
    def p(texto): return Paragraph(escape(str(texto)),estilos['CeldaBanco'])
    def tabla_pdf(headers, rows, anchos):
        t=PDFTable([[p(h) for h in headers]]+rows,colWidths=anchos,repeatRows=1,hAlign='LEFT')
        t.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),'Informe'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DCE8F0')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F3F6F8')]),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ALIGN',(1,1),(-1,-1),'RIGHT'),('FONTSIZE',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),3),('TOPPADDING',(0,0),(-1,-1),3),('LINEBELOW',(0,0),(-1,0),0.6,colors.HexColor('#17324D'))]))
        return t
    sistema=tabla[(tabla.fecha==mes)&(tabla.tipo_entidad=='sistema')]
    seleccion=tabla[(tabla.fecha==mes)&(tabla.tipo_entidad=='banco')&tabla.banco.isin(bancos)].sort_values('banco')
    grupos=[list(INDICADORES)[:4],list(INDICADORES)[4:]]
    for indice,grupo in enumerate(grupos):
        if indice:
            contenido.extend([PageBreak(),Paragraph('Rentabilidad, utilidades y participación',estilos['Heading1']),Paragraph(f'Cierre: {mes}. ROE y ROA anualizados, después de impuestos. Utilidad acumulada desde enero.',estilos['Normal']),Spacer(1,8)])
        contenido.append(Paragraph('Sistema bancario completo',estilos['Heading2']))
        if len(sistema):
            fila=sistema.iloc[0]
            contenido.append(tabla_pdf([INDICADORES[c] for c in grupo],[[numero(fila[c],0 if c.endswith('mmclp') else 2) for c in grupo]],[195]*len(grupo)))
        contenido.extend([Spacer(1,8),Paragraph('Bancos seleccionados',estilos['Heading2'])])
        filas=[[p(r.banco)]+[numero(r[c],0 if c.endswith('mmclp') else 2) for c in grupo] for _,r in seleccion.iterrows()]
        if filas:
            contenido.append(tabla_pdf(['Banco',*[INDICADORES[c] for c in grupo]],filas,[230,140,120,145,145]))
        else:
            contenido.append(Paragraph('No hay bancos seleccionados con datos en este período.',estilos['Normal']))
    if referencia and bancos:
        contenido.extend([PageBreak(),Paragraph('Cambios por banco',estilos['Heading1']),Paragraph(escape(INDICADORES[indicador])+f' | {referencia} a {mes}',estilos['Normal']),Spacer(1,10)])
        cambios=comparar(tabla,mes,referencia,indicador,bancos)
        monto=indicador.endswith('mmclp')
        cols=['referencia','actual','cambio']+(['variacion_pct'] if monto else [])
        filas=[[p(r.banco)]+[numero(r[c],3 if c=='variacion_pct' else 2) for c in cols] for _,r in cambios.iterrows()]
        headers=['Banco',referencia,mes,'Cambio (MM CLP)' if monto else 'Cambio (puntos porcentuales)']+(['Variación nominal (%)'] if monto else [])
        anchos=[230]+([137.5]*4 if monto else [183.33]*3)
        contenido.append(tabla_pdf(headers,filas,anchos))
    contenido.extend([PageBreak(),Paragraph('Definiciones y fuentes',estilos['Heading1'])])
    for nota in NOTAS:
        contenido.extend([Paragraph(escape(nota),estilos['Normal']),Spacer(1,9)])
    contenido.append(Paragraph('Archivos utilizados',estilos['Heading2']))
    fechas={mes,referencia} if referencia else {mes}
    for nombre in sorted(tabla[tabla.fecha.isin(fechas)].archivo_fuente.dropna().unique()):
        contenido.append(Paragraph(escape(nombre),estilos['Normal']))
    contenido.extend([Spacer(1,12),Paragraph(FUENTE,estilos['Normal'])])
    def pie(canvas,documento):
        canvas.setFont('Informe',8)
        canvas.setFillColor(colors.HexColor('#546576'))
        canvas.drawString(30,18,'Fuente: CMF. Informe generado a partir de los datos guardados y filtros seleccionados.')
        canvas.drawRightString(810,18,f'Página {documento.page}')
    doc.build(contenido,onFirstPage=pie,onLaterPages=pie)
    return buf.getvalue()