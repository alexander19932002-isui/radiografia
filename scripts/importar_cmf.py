
"""Importa el formato mensual CMF validado con el reporte de agosto de 2026."""
from pathlib import Path
import calendar
import re
import sqlite3
import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
MESES = dict(zip('ENERO FEBRERO MARZO ABRIL MAYO JUNIO JULIO AGOSTO SEPTIEMBRE OCTUBRE NOVIEMBRE DICIEMBRE'.split(), range(1, 13)))
MAPA = {
    'colocaciones_mmclp': ('Est. Situación Financ. Bancos', 4),
    'colocaciones_costo_amortizado_mmclp': ('Est. Situación Financ. Bancos', 5),
    'morosidad_90_pct': ('Ind. Mora y Deteriorada', 3),
    'provisiones_clasificacion_pct': ('Ind. R. Crédito Provisiones', 3),
    'eficiencia_gastos_ingresos_pct': ('Ind. de Rentab. y Eficiencia', 13),
    'roe_pct': ('Ind. de Rentab. y Eficiencia', 4),
    'roa_pct': ('Ind. de Rentab. y Eficiencia', 6),
    'utilidad_acumulada_mmclp': ('Est. del Resultado Bancos', 12),
}

def normalizar(texto):
    return re.sub(r'\s*\(\d+\)\s*$', '', str(texto)).strip()

def extraer(archivo, nombre):
    libro = openpyxl.load_workbook(archivo, data_only=True, read_only=True)
    try:
        hojas = {s.title.strip(): s for s in libro}
        titulo = ' '.join(str(c) for row in hojas['Índice'].iter_rows(values_only=True) for c in row if c)
        fecha = re.search(r'A\s+(' + '|'.join(MESES) + r')\s+DE\s+(\d{4})', titulo.upper())
        if not fecha:
            raise ValueError('No se pudo identificar el mes en la hoja Índice.')
        mes, anio = MESES[fecha[1]], int(fecha[2])
        periodo = f'{anio}-{mes:02d}-{calendar.monthrange(anio, mes)[1]}'
        datos = {}
        trazas = []
        for indicador, (nombre_hoja, columna) in MAPA.items():
            if nombre_hoja not in hojas:
                raise ValueError(f'Falta la hoja {nombre_hoja}. Revisa el formato del reporte.')
            hoja = hojas[nombre_hoja]
            # Validar la columna antes de extraer: cambios de formato requieren revisión.
            cabecera = ' '.join(str(hoja.cell(r, columna).value or '') for r in range(7, 12)).lower()
            esperado = {'colocaciones_mmclp':'total', 'colocaciones_costo_amortizado_mmclp':'costo amortizado', 'morosidad_90_pct':'morosidad', 'provisiones_clasificacion_pct':'clasificación', 'eficiencia_gastos_ingresos_pct':'ingresos', 'roe_pct':'(roe)', 'roa_pct':'(roa)', 'utilidad_acumulada_mmclp':'resultado del ejercicio consolidado'}[indicador]
            if esperado not in cabecera:
                raise ValueError(f'Cambió la cabecera de {nombre_hoja}, columna {columna}. No se importó el archivo.')
            encontrados = set()
            for fila in hoja.iter_rows():
                banco = normalizar(fila[1].value) if len(fila) > 1 else ''
                es_sistema = banco == 'Sistema Bancario'
                es_banco = banco.startswith(('Banco ', 'Bank of ', 'China Construction ', 'HSBC ', 'Jp Morgan ', 'Scotiabank ', 'Tanner Banco '))
                valor = fila[columna - 1].value
                if not (es_sistema or es_banco) or not isinstance(valor, (float, int)):
                    continue
                if banco in encontrados:
                    raise ValueError(f'Banco duplicado en {nombre_hoja}: {banco}')
                encontrados.add(banco)
                registro = datos.setdefault(banco, {'fecha': periodo, 'banco': banco, 'tipo_entidad': 'sistema' if es_sistema else 'banco', 'alcance': 'consolidado', 'archivo_fuente': nombre})
                registro[indicador] = valor
                trazas.append({'fecha': periodo, 'banco': banco, 'indicador': indicador, 'valor': valor, 'unidad': 'MM CLP' if indicador.endswith('mmclp') else 'porcentaje (0-100)', 'archivo': nombre, 'hoja': hoja.title, 'celda': fila[columna - 1].coordinate})
            if not encontrados:
                raise ValueError(f'No se encontraron cifras válidas en {nombre_hoja}.')
        tabla = pd.DataFrame(datos.values())
        if tabla[list(MAPA)].isna().any().any():
            raise ValueError('Faltan indicadores para una o más entidades. Revisar antes de importar.')
        return tabla.sort_values(['tipo_entidad', 'banco']), pd.DataFrame(trazas)
    finally:
        libro.close()

def main():
    import sys
    sys.path.insert(0, str(ROOT))
    from scripts.almacenamiento import preparar, guardar, leer, conexion
    archivos = []
    for archivo in sorted((ROOT / 'datos/originales').glob('*.xlsx')):
        archivos.append((archivo.name, archivo.read_bytes()))
    if not archivos:
        raise SystemExit('Coloca al menos un reporte XLSX de la CMF en datos/originales.')
    lote = preparar(archivos)
    existentes = leer(ROOT)
    repetidos = sorted(set(existentes.fecha) & {item[2].fecha.iloc[0] for item in lote})
    reemplazar = '--reemplazar' in sys.argv
    # No descartar cargas anteriores del panel. Los meses repetidos exigen opción explícita.
    if not reemplazar:
        lote = [item for item in lote if item[2].fecha.iloc[0] not in repetidos]
    if lote:
        guardar(ROOT, lote, reemplazar=reemplazar)
    tabla = leer(ROOT)
    salida = ROOT / 'datos/procesados'
    salida.mkdir(exist_ok=True, parents=True)
    tabla.to_csv(salida / 'bancos_cmf.csv', index=False, encoding='utf-8-sig')
    with conexion(ROOT) as con:
        pd.read_sql_query('SELECT * FROM trazabilidad ORDER BY fecha,banco,indicador', con).to_csv(salida / 'trazabilidad.csv', index=False, encoding='utf-8-sig')
    print(f'Historial: {len(tabla)} registros. Meses: {tabla.fecha.nunique()}. Reportes incorporados: {len(lote)}.')

if __name__ == '__main__':
    main()